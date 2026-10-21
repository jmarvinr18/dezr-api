"""Test doubles and file builders for the document ingestion tests."""

import io
import random

from docx import Document as new_docx

from app.constants import EMBEDDING_DIMENSIONS
from app.services.ingestion import Embedder, EmbeddingUnavailable


class FakeEmbedder(Embedder):
    """Deterministic vectors; records what it was asked to embed."""

    model_id = "fake-embed-v1"

    def __init__(self, fail: bool = False, on_embed=None):
        self.fail = fail
        self.on_embed = on_embed
        self.calls: list[list[str]] = []

    def embed(self, texts):
        self.calls.append(list(texts))
        if self.on_embed:
            self.on_embed()
        if self.fail:
            raise EmbeddingUnavailable("embedding service down")
        return [self._vector(text) for text in texts]

    @staticmethod
    def _vector(text: str) -> list[float]:
        rng = random.Random(text)
        return [rng.uniform(-1, 1) for _ in range(EMBEDDING_DIMENSIONS)]


# ---------- PDF ----------


def _escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages: list[dict]) -> bytes:
    """A minimal PDF with Helvetica text. Each page: {"lines": [(text, x, y_from_top, size, bold)],
    "rules": [(x0, top0, x1, top1)]} where rules are straight lines (for drawing tables)."""
    height = 792
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    add(b"<< /Type /Catalog /Pages 2 0 R >>")
    add(b"")  # Pages, filled in below
    regular = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    bold = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>")
    page_ids = []
    for page in pages:
        ops = []
        for x0, top0, x1, top1 in page.get("rules", []):
            ops.append(f"0.5 w {x0} {height - top0} m {x1} {height - top1} l S")
        for text, x, top, size, is_bold in page.get("lines", []):
            font = "F2" if is_bold else "F1"
            ops.append(f"BT /{font} {size} Tf {x} {height - top} Td ({_escape(text)}) Tj ET")
        stream = "\n".join(ops).encode("cp1252")
        content = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        page_ids.append(
            add(
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 {height}] "
                f"/Resources << /Font << /F1 {regular} 0 R /F2 {bold} 0 R >> >> /Contents {content} 0 R >>".encode()
            )
        )
    kids = " ".join(f"{i} 0 R" for i in page_ids)
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def table_rules(x0: float, top: float, widths: list[float], row_height: float, rows: int):
    """Ruling lines for a grid table (what pdfplumber detects tables by)."""
    x1 = x0 + sum(widths)
    rules = [(x0, top + r * row_height, x1, top + r * row_height) for r in range(rows + 1)]
    x = x0
    for width in [0, *widths]:
        x += width
        rules.append((x, top, x, top + rows * row_height))
    return rules


# ---------- Word ----------


def make_docx() -> bytes:
    """A Word file with a title, headings, a list, a table and body text."""
    document = new_docx()
    document.add_heading("Engagement Guide", 0)
    document.add_heading("Fees", 1)
    document.add_paragraph("Our fees depend on the matter. We agree them in writing before work starts.")
    document.add_paragraph("Hourly work", style="List Bullet")
    document.add_paragraph("Fixed fees for incorporations", style="List Bullet")
    document.add_heading("Retainers", 2)
    document.add_paragraph("Retainers are billed monthly in advance.")
    table = document.add_table(rows=3, cols=2)
    for r, row in enumerate([("Service", "Fee"), ("Incorporation", "PHP 50,000"), ("Visa", "PHP 30,000")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.add_heading("Contact", 1)
    document.add_paragraph("Email admin@deszr.com for an engagement letter.")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
